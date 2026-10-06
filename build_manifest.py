import os, json, re, shutil, sys
import openpyxl

REPO = "/home/tina/Documents/lerobot_pi05"
SITE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(SITE, "data")
XLSX = "/home/tina/Documents/ieeeconf-git/baselines_lpips_jf.xlsx"

TASKS = ["Place Mug", "Shelf Book", "Pour Liquid", "Pick Shoe"]
TASK_KEY = {"Shelf Book": "book_shelving", "Place Mug": "place_mug", "Pour Liquid": "pouring", "Pick Shoe": "pick_shoe"}
CAMERAS = ["Stationary", "Wrist"]
CAM_KEY = {"Stationary": "stationary", "Wrist": "wrist"}

# Display name (paper naming, Table IV order) -> row name in baselines_lpips_jf.xlsx
METHODS = {
    "Raw Render": "Raw Render",
    "Classical Color Alignment": "Classical Color Alignment",
    "Pix2Pix": "Pix2Pix",
    "Pix2Pix-DINO": "Pix2Pix-DINO",
    "Pix2Pix-DINO w/o Pixel": "Pix2Pix-DINO w/o Pixel",
    "Pix2Pix GAN-Only": "Pix2Pix GAN-Only",
    "Turbo-Reconstruction": "Paired-Default",
    "GAN-Only": "GAN-Only",
    "DINO-Only": "DINO-Only",
    "STRIPE+L2": "DINO-Align+L2",
    "STRIPE+LPIPS": "DINO-Align+LPIPS",
    "STRIPE+Paired": "DINO-Align+Paired",
    "STRIPE (Ours)": "DINO-Align (Ours)",
}
METHOD_ORDER = list(METHODS)

def dirs_for(method, task_key, cam):
    # cam: 'stationary' or 'wrist'
    if method == "Raw Render" or method == "Classical Color Alignment":
        return f"outputs/color_calib_{cam}_{short(task_key)}/results"
    if method == "Pix2Pix":
        return f"outputs/pix2pix_{cam}_{short(task_key)}/results"
    if method == "Pix2Pix-DINO":
        return "outputs/pix2pix_dino_all_tasks/results"
    if method == "Pix2Pix-DINO w/o Pixel":
        return "outputs/pix2pix_dino_all_tasks_no_l1/results"
    if method == "Pix2Pix GAN-Only":
        return "outputs/pix2pix_all_tasks_gan_only/results"
    if method == "Turbo-Reconstruction":
        return "outputs/turbo_sim2real_all_tasks_paired_default/results"
    if method == "GAN-Only":
        return "outputs/turbo_sim2real_all_tasks_gan_only/results"
    if method == "DINO-Only":
        return "outputs/turbo_sim2real_all_tasks_dino_only/results"
    if method == "STRIPE+L2":
        return "outputs/turbo_sim2real_all_tasks_ours_l2/results"
    if method == "STRIPE+LPIPS":
        return "outputs/turbo_sim2real_all_tasks_ours_lpips/results"
    if method == "STRIPE+Paired":
        return "outputs/turbo_sim2real_all_tasks_ours_paired/results"
    if method == "STRIPE (Ours)":
        # joint all-tasks checkpoint, the model reported in the paper
        return "outputs/turbo_sim2real_all_tasks/results"
    raise ValueError(method)

def short(task_key):
    return {"book_shelving": "book", "place_mug": "mug", "pouring": "pouring", "pick_shoe": "shoe"}[task_key]

def short2(task_key):
    return {"book_shelving": "book", "pouring": "pouring", "pick_shoe": "shoe"}[task_key]

def find_result_subdir(results_root, task_key, cam):
    if not os.path.isdir(os.path.join(REPO, results_root)):
        return None
    for name in os.listdir(os.path.join(REPO, results_root)):
        if name.startswith(f"{task_key}_{cam}_") or name == f"{task_key}_{cam}":
            full = os.path.join(REPO, results_root, name, "test_latest", "images")
            if os.path.isdir(full):
                return full
    return None

def list_triples(images_dir, need_fake=True):
    idxs = {}
    for f in os.listdir(images_dir):
        m = re.match(r"^(\d+)_(real_A|real_B|fake_B)\.png$", f)
        if not m:
            continue
        idx, kind = m.group(1), m.group(2)
        idxs.setdefault(idx, {})[kind] = f
    out = []
    for idx, d in sorted(idxs.items()):
        need = {"real_A", "real_B"} | ({"fake_B"} if need_fake else set())
        if need.issubset(d.keys()):
            out.append(idx)
    return out, idxs

def rel_symlink(dst, src_abs):
    # copies the file (name kept for history); the published site needs real files
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    if os.path.lexists(dst):
        os.remove(dst)
    shutil.copy2(src_abs, dst)

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["All values (tidy)"]
    metrics = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        method, task, cam, lpips, jf, j, f, as_ = row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7]
        metrics[(method, task, cam)] = {"lpips": lpips, "jf": jf, "j": j, "f": f, "as": as_}

    entries = []
    missing = []
    for method in METHOD_ORDER:
        for task in TASKS:
            task_key = TASK_KEY[task]
            for cam in CAMERAS:
                cam_key = CAM_KEY[cam]
                results_root = dirs_for(method, task_key, cam_key)
                images_dir = find_result_subdir(results_root, task_key, cam_key)
                if images_dir is None:
                    missing.append((method, task, cam, results_root))
                    continue
                need_fake = method != "Raw Render"
                idxs, idxmap = list_triples(images_dir, need_fake=need_fake)
                if not idxs:
                    missing.append((method, task, cam, images_dir))
                    continue
                mkey = re.sub(r"[^a-z0-9]+", "_", method.lower()).strip("_")
                out_samples = []
                for idx in idxs:
                    d = idxmap[idx]
                    dest_dir = os.path.join(OUT_DIR, mkey, f"{task_key}_{cam_key}")
                    a_src = os.path.join(images_dir, d["real_A"])
                    b_src = os.path.join(images_dir, d["real_B"])
                    a_dst = os.path.join(dest_dir, f"{idx}_real_A.png")
                    b_dst = os.path.join(dest_dir, f"{idx}_real_B.png")
                    rel_symlink(a_dst, a_src)
                    rel_symlink(b_dst, b_src)
                    if need_fake:
                        f_src = os.path.join(images_dir, d["fake_B"])
                        f_dst = os.path.join(dest_dir, f"{idx}_fake_B.png")
                        rel_symlink(f_dst, f_src)
                    else:
                        f_dst = a_dst  # raw render: model output == input render
                    out_samples.append({
                        "idx": idx,
                        "real_A": os.path.relpath(a_dst, SITE),
                        "real_B": os.path.relpath(b_dst, SITE),
                        "fake_B": os.path.relpath(f_dst, SITE),
                    })
                m = metrics.get((METHODS[method], task, cam), {})
                entries.append({
                    "method": method,
                    "method_key": mkey,
                    "task": task,
                    "task_key": task_key,
                    "camera": cam,
                    "camera_key": cam_key,
                    "lpips": m.get("lpips"),
                    "jf": m.get("jf"),
                    "j": m.get("j"),
                    "f": m.get("f"),
                    "as": m.get("as"),
                    "n_samples": len(out_samples),
                    "samples": out_samples,
                })

    ws_mean = wb["Mean summary"]
    means = {}
    for row in ws_mean.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        means[row[0]] = {"lpips": row[1], "jf": row[2], "j": row[3], "f": row[4], "as": row[5], "jf_artifact": row[6]}

    manifest = {
        "methods": METHOD_ORDER,
        "method_keys": {m: re.sub(r"[^a-z0-9]+", "_", m.lower()).strip("_") for m in METHOD_ORDER},
        "tasks": TASKS,
        "cameras": CAMERAS,
        "means": {d: means.get(x, {}) for d, x in METHODS.items()},
        "entries": entries,
    }
    with open(os.path.join(SITE, "manifest.json"), "w") as f:
        json.dump(manifest, f)

    print(f"entries built: {len(entries)} / {len(METHOD_ORDER)*len(TASKS)*len(CAMERAS)}")
    print(f"total samples: {sum(e['n_samples'] for e in entries)}")
    if missing:
        print(f"MISSING ({len(missing)}):")
        for m in missing:
            print(" ", m)

if __name__ == "__main__":
    main()
